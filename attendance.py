import tkinter as tk
from tkinter import messagebox
import openpyxl
import os
from datetime import date

FILE_NAME = "attendance_records.xlsx"

# ---------------- CREATE EXCEL FILE ----------------

if not os.path.exists(FILE_NAME):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    ws.append([
        "Record ID",
        "Student ID",
        "Student Name",
        "Class",
        "Date",
        "Subject",
        "Attendance"
    ])

    wb.save(FILE_NAME)
    wb.close()


# ---------------- ADD RECORD ----------------

def add_record():

    student_id = student_id_entry.get().strip()
    student_name = name_entry.get().strip()
    student_class = class_entry.get().strip()
    attendance_date = date_entry.get().strip()
    subject = subject_entry.get().strip()
    attendance = attendance_var.get()

    if (student_id == "" or
        student_name == "" or
        student_class == "" or
        attendance_date == "" or
        subject == ""):

        messagebox.showwarning(
            "Warning",
            "Please fill all fields."
        )
        return

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    record_id = ws.max_row

    ws.append([
        record_id,
        student_id,
        student_name,
        student_class,
        attendance_date,
        subject,
        attendance
    ])

    wb.save(FILE_NAME)
    wb.close()

    messagebox.showinfo(
        "Success",
        "Attendance added successfully!"
    )

    clear_form()


# ---------------- SHOW RECORDS ----------------

def show_records():

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    records = ""

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        records += (
            "Record ID: " + str(row[0]) +
            "\nStudent ID: " + str(row[1]) +
            "\nStudent Name: " + str(row[2]) +
            "\nClass: " + str(row[3]) +
            "\nDate: " + str(row[4]) +
            "\nSubject: " + str(row[5]) +
            "\nAttendance: " + str(row[6]) +
            "\n--------------------------\n"
        )

    wb.close()

    if records == "":
        records = "No attendance records found."

    messagebox.showinfo(
        "Attendance Records",
        records
    )


# ---------------- SEARCH RECORD ----------------

def search_record():

    search_value = search_entry.get().strip().lower()

    if search_value == "":
        messagebox.showwarning(
            "Search",
            "Enter Student ID or Student Name."
        )
        return

    wb = openpyxl.load_workbook(FILE_NAME)
    ws = wb.active

    records = ""

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        student_id = str(row[1]).lower()
        student_name = str(row[2]).lower()

        if (search_value in student_id or
            search_value in student_name):

            records += (
                "Record ID: " + str(row[0]) +
                "\nStudent ID: " + str(row[1]) +
                "\nStudent Name: " + str(row[2]) +
                "\nClass: " + str(row[3]) +
                "\nDate: " + str(row[4]) +
                "\nSubject: " + str(row[5]) +
                "\nAttendance: " + str(row[6]) +
                "\n--------------------------\n"
            )

    wb.close()

    if records == "":
        records = "No matching record found."

    messagebox.showinfo(
        "Search Result",
        records
    )


# ---------------- CLEAR FORM ----------------

def clear_form():

    student_id_entry.delete(0, tk.END)
    name_entry.delete(0, tk.END)
    class_entry.delete(0, tk.END)
    subject_entry.delete(0, tk.END)
    search_entry.delete(0, tk.END)

    date_entry.delete(0, tk.END)
    date_entry.insert(0, str(date.today()))

    attendance_var.set("Present")


# ---------------- MAIN WINDOW ----------------

root = tk.Tk()

root.title(
    "Students Attendance Management System"
)

root.geometry("400x750")

root.configure(
    bg="#EAF4F4"
)


# ---------------- TITLE ----------------

title = tk.Label(
    root,
    text="Students Attendance\nManagement System",
    font=("Arial", 17, "bold"),
    bg="#145A32",
    fg="white",
    pady=15
)

title.pack(fill="x")


# ---------------- FORM ----------------

form = tk.Frame(
    root,
    bg="#EAF4F4"
)

form.pack(
    fill="both",
    expand=True,
    padx=20,
    pady=12
)


# Student ID

tk.Label(
    form,
    text="Student ID",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")

student_id_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

student_id_entry.pack(
    fill="x",
    pady=(2, 7)
)


# Student Name

tk.Label(
    form,
    text="Student Name",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")

name_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

name_entry.pack(
    fill="x",
    pady=(2, 7)
)


# Class

tk.Label(
    form,
    text="Class",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")

class_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

class_entry.pack(
    fill="x",
    pady=(2, 7)
)


# Date

tk.Label(
    form,
    text="Date",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")

date_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

date_entry.pack(
    fill="x",
    pady=(2, 7)
)

date_entry.insert(
    0,
    str(date.today())
)


# Subject

tk.Label(
    form,
    text="Subject",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")

subject_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

subject_entry.pack(
    fill="x",
    pady=(2, 7)
)


# Attendance

tk.Label(
    form,
    text="Attendance",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(anchor="w")


attendance_var = tk.StringVar()

attendance_var.set("Present")


attendance_menu = tk.OptionMenu(
    form,
    attendance_var,
    "Present",
    "Absent"
)

attendance_menu.config(
    font=("Arial", 12)
)

attendance_menu.pack(
    fill="x",
    pady=(2, 10)
)


# ---------------- ADD BUTTON ----------------

tk.Button(
    form,
    text="ADD ATTENDANCE",
    font=("Arial", 11, "bold"),
    bg="#27AE60",
    fg="white",
    command=add_record
).pack(
    fill="x",
    pady=3
)


# ---------------- SEARCH ----------------

tk.Label(
    form,
    text="Search Student",
    font=("Arial", 12, "bold"),
    bg="#EAF4F4"
).pack(
    anchor="w",
    pady=(8, 0)
)


search_entry = tk.Entry(
    form,
    font=("Arial", 12)
)

search_entry.pack(
    fill="x",
    pady=3
)


tk.Button(
    form,
    text="SEARCH",
    font=("Arial", 11, "bold"),
    bg="#8E44AD",
    fg="white",
    command=search_record
).pack(
    fill="x",
    pady=3
)


# ---------------- SHOW RECORDS ----------------

tk.Button(
    form,
    text="SHOW ALL RECORDS",
    font=("Arial", 11, "bold"),
    bg="#2980B9",
    fg="white",
    command=show_records
).pack(
    fill="x",
    pady=3
)


# ---------------- CLEAR ----------------

tk.Button(
    form,
    text="CLEAR",
    font=("Arial", 11, "bold"),
    command=clear_form
).pack(
    fill="x",
    pady=3
)


# ---------------- START PROGRAM ----------------

root.mainloop()